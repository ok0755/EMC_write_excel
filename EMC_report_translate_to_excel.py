#coding=utf-8
#author:Zhoubin
#program language:Python2.7
'''
功能：将本厂EMC报告word格式转置成excel格式，过滤保留margin dB最低值，及对应文件名、频率，生成EMC.xlsx文档
使用方法：任意数量本厂格式EMC报告与本程序放任意同一文件夹，执行本程序，同文件夹内生成EMC.xlsx
'''
import docx                                                          # 导入word读写库
import os
import xlsxwriter                                                    # 导入excel处理库
import win32api                                                      # windows API接口

# 获取.docx文件
def get_file():
    for f in os.listdir(file_dir):
        if 'docx' in f:
            yield f
# 主程序
def main():
    xl_book = xlsxwriter.Workbook('EMC.xlsx')                          # 同目录下新建EMC.xlsx
    xl = xl_book.add_worksheet('sheet1')                               # sheet_name=sheet1
    xl.write(0,0,u'文件名')                                            # 写入excel首行
    xl.write(0,1,u'频率')
    xl.write(0,2,u'Margin QP')
    j = 1
    for ff in get_file():                                                 # 遍历word文档
        file_path = os.path.join(file_dir,ff)
        doc_file = docx.Document(file_path)
        try:                                                              # 忽略非EMC报告格式word文档
            table = doc_file.tables[0]                                    # word文档第一个表格
            length=len(table.rows)                                        # 数据总行数
            ar = [float(table.cell(i,13).text) for i in range(3,length)]  # dB值数列
            index = ar.index(min(ar))
            cell_1 = float(table.cell(index+3,1).text)
            cell_2 = min(ar)                                               # 取最小dB值
            xl.write(j,0,ff.decode('gb18030','ignore'))                    # 写入excel
            xl.write(j,1,cell_1)
            xl.write(j,2,min(ar))
            j+=1
        except:
            pass

    xl.set_column('A:A',25)                                                # 设置excel列宽
    xl.set_column('B:B',15)
    xl.set_column('C:C',15)
    xl.set_column('D:D',15)
    xl_book.close()
    path=os.path.join(file_dir,'EMC.xlsx')
    win32api.ShellExecute(0,'open',path,'','',1)                           # 调用windows系统命令，双击打开EMC.xlsx

if __name__=='__main__':
    file_dir=os.getcwd()   # 当前目录
    main()
