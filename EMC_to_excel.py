#coding=utf-8
import docx
import os
import xlsxwriter

def get_file():
    for f in os.listdir(file_dir):
        if 'docx' in f:
            yield f
        #yield f

if __name__=='__main__':
    file_dir = os.getcwd()
    xl_book = xlsxwriter.Workbook('EMC.xlsx')
    xl = xl_book.add_worksheet('sheet1')
    xl.write(0,1,u'频率')
    xl.write(0,2,u'dB余量')

    j = 1
    for ff in get_file():
        file_path = os.path.join(file_dir,ff)
        doc_file = docx.Document(file_path)
        table = doc_file.tables[0]
        cols=table.columns
        for i in range(3,len(table.rows)):
            xl.write(j,0,ff.decode('gb18030','ignore'))
            xl.write(j,1,table.cell(i,1).text)
            xl.write(j,2,table.cell(i,13).text)
            j+=1
    xl_book.close()